"""Links and skill groups in resumes: parsing rules, Word hyperlinks and the text ATS sees."""

import io
import unittest

from docx import Document

from app.schemas.resume import ResumeEntry, StructuredResume
from app.services.resume_docx import build_resume_docx, get_template
from app.services.resume_sections import render_resume_text
from app.services.rich_text import group_skills, plain_links, safe_url, short_url, split_links

RESUME = StructuredResume(
    name="Asha Rao",
    email="asha@example.com",
    phone="+91 98765 43210",
    links=["github.com/asha-dev", "linkedin.com/in/asha"],
    summary="Builds web apps. Portfolio: [asha.dev](https://asha.dev).",
    skills=["Languages: Python", "Java", "Tools: Git", "Docker"],
    projects=[
        ResumeEntry(
            title="Event Portal",
            subtitle="React, FastAPI",
            date="2025",
            link="github.com/asha-dev/portal",
            bullets=["Built a portal used by 300+ students ([demo](https://portal.example.com))."],
        )
    ],
    certifications=["[AWS Cloud Practitioner](https://aws.amazon.com/certification/)", "[Bad](javascript:alert(1))"],
)


def _links(content: bytes) -> list[str]:
    document = Document(io.BytesIO(content))
    return [rel.target_ref for rel in document.part.rels.values() if rel.reltype.endswith("/hyperlink")]


class SafeUrlTests(unittest.TestCase):
    def test_adds_scheme_and_handles_email(self):
        self.assertEqual(safe_url("github.com/asha"), "https://github.com/asha")
        self.assertEqual(safe_url("https://x.dev/a"), "https://x.dev/a")
        self.assertEqual(safe_url("asha@example.com"), "mailto:asha@example.com")

    def test_refuses_unsafe_or_non_urls(self):
        for value in ["javascript:alert(1)", "data:text/html,hi", "file:///etc/passwd", "not a url", "Bengaluru", ""]:
            with self.subTest(value=value):
                self.assertIsNone(safe_url(value))

    def test_short_url(self):
        self.assertEqual(short_url("https://www.github.com/asha/"), "github.com/asha")


class InlineLinkTests(unittest.TestCase):
    def test_split_and_plain_text(self):
        parts = split_links("See [my demo](demo.example.com) now")
        self.assertEqual(parts, [("See ", None), ("my demo", "https://demo.example.com"), (" now", None)])
        self.assertEqual(
            plain_links("See [my demo](https://demo.example.com) now"), "See my demo (demo.example.com) now"
        )

    def test_unsafe_link_stays_plain_text(self):
        self.assertEqual(split_links("[x](javascript:alert(1))")[0][1], None)


class SkillGroupTests(unittest.TestCase):
    def test_labels_start_groups(self):
        self.assertEqual(
            group_skills(["Python", "Languages: Java", "C++", "Tools:", "Git", " "]),
            [("", ["Python"]), ("Languages", ["Java", "C++"]), ("Tools", ["Git"])],
        )

    def test_plain_skills_are_one_group(self):
        self.assertEqual(group_skills(["Python", "SQL"]), [("", ["Python", "SQL"])])


class OutputTests(unittest.TestCase):
    def test_ats_text_shows_link_addresses(self):
        text = render_resume_text(RESUME)
        self.assertIn("github.com/asha-dev/portal", text)
        self.assertIn("AWS Cloud Practitioner (aws.amazon.com/certification)", text)
        self.assertIn("demo (portal.example.com)", text)
        self.assertNotIn("](", text.replace("[Bad](javascript:alert(1))", ""))

    def test_word_file_has_real_hyperlinks(self):
        links = _links(build_resume_docx(RESUME, "tech-classic"))
        for url in [
            "mailto:asha@example.com",
            "tel:+919876543210",
            "https://github.com/asha-dev",
            "https://linkedin.com/in/asha",
            "https://github.com/asha-dev/portal",
            "https://portal.example.com",
            "https://asha.dev",
            "https://aws.amazon.com/certification/",
        ]:
            with self.subTest(url=url):
                self.assertIn(url, links)
        self.assertFalse(any(link.startswith("javascript") for link in links))

    def test_word_text_is_readable(self):
        document = Document(io.BytesIO(build_resume_docx(RESUME, "tech-classic")))
        text = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("Languages: Python, Java", text)
        self.assertIn("Tools: Git, Docker", text)
        self.assertIn("Event Portal  github.com/asha-dev/portal\t2025", text)
        self.assertIn("AWS Cloud Practitioner", text)
        self.assertNotIn("](", text.replace("[Bad](javascript:alert(1))", ""))

    def test_tech_templates_exist_and_icon_templates_split_contact(self):
        for template_id in ("tech-classic", "tech-modern", "tech-compact"):
            self.assertEqual(get_template(template_id)["id"], template_id)
        document = Document(io.BytesIO(build_resume_docx(RESUME, "tech-classic")))
        lines = [p.text for p in document.paragraphs[:4]]
        self.assertIn("asha@example.com | +91 98765 43210", lines)
        self.assertIn("github.com/asha-dev | linkedin.com/in/asha", lines)


if __name__ == "__main__":
    unittest.main()
